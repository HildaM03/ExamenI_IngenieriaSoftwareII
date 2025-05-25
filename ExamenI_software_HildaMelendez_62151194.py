from datetime import datetime, timedelta
import os
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from abc import ABC, abstractmethod

# Factory Method Pattern
class EntityFactory(ABC):
    @abstractmethod
    def create_entity(self, *args, **kwargs):
        pass

class ServicioFactory(EntityFactory):
    def create_entity(self, id, nombre, duracion, precio):
        return Database.Servicio(id, nombre, duracion, precio)

class ClienteFactory(EntityFactory):
    def create_entity(self, id, nombre, email, telefono):
        return Database.Cliente(id, nombre, email, telefono)

class CitaFactory(EntityFactory):
    def create_entity(self, id, cliente, servicio, fecha_hora):
        return Database.Cita(id, cliente, servicio, fecha_hora)

# Command Pattern
class Command(ABC):
    @abstractmethod
    def execute(self):
        pass

class AgendarCitaCommand(Command):
    def __init__(self, db, cliente, servicio, fecha_hora):
        self.db = db
        self.cliente = cliente
        self.servicio = servicio
        self.fecha_hora = fecha_hora
    
    def execute(self):
        if not self.db.verificar_disponibilidad(self.servicio, self.fecha_hora):
            return False, "No hay disponibilidad para ese horario"
        
        cita = self.db.agregar_cita(self.cliente, self.servicio, self.fecha_hora)
        return True, cita

class CancelarCitaCommand(Command):
    def __init__(self, db, cita_id):
        self.db = db
        self.cita_id = cita_id
    
    def execute(self):
        for cita in self.db.citas:
            if cita.id == self.cita_id:
                cita.estado = "Cancelada"
                self.db.save_data()
                return True, "Cita cancelada correctamente"
        return False, "No se encontró la cita"

class BuscarCitasCommand(Command):
    def __init__(self, db, cliente_id=None, fecha=None):
        self.db = db
        self.cliente_id = cliente_id
        self.fecha = fecha
    
    def execute(self):
        if self.cliente_id:
            return True, [c for c in self.db.citas 
                         if c.cliente.id == self.cliente_id 
                         and c.estado == "Confirmada"]
        elif self.fecha:
            return True, [c for c in self.db.citas 
                         if c.fecha_hora.date() == self.fecha.date() 
                         and c.estado == "Confirmada"]
        return False, "Parámetros de búsqueda no válidos"

# Decorator Pattern
class CitaDecorator(ABC):
    def __init__(self, cita):
        self.cita = cita
    
    @abstractmethod
    def get_details(self):
        pass

class CitaConNotificacion(CitaDecorator):
    def get_details(self):
        details = self.cita.get_details()
        details += "\nNotificación: Se enviará recordatorio 24 horas antes"
        return details

class CitaConDescuento(CitaDecorator):
    def __init__(self, cita, descuento):
        super().__init__(cita)
        self.descuento = descuento
    
    def get_details(self):
        details = self.cita.get_details()
        precio_con_descuento = self.cita.servicio.precio * (1 - self.descuento/100)
        details += f"\nDescuento: {self.descuento}% - Precio final: ${precio_con_descuento:.2f}"
        return details

class Database:
    _instance = None
    
    def __new__(cls):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
            cls._instance._init_db()
        return cls._instance
    
    def _init_db(self):
        servicio_factory = ServicioFactory()
        self.servicios = [
            servicio_factory.create_entity(1, "Corte de Cabello", 30, 150.00),
            servicio_factory.create_entity(2, "Manicure", 45, 200.00),
            servicio_factory.create_entity(3, "Pedicure", 60, 250.00),
            servicio_factory.create_entity(4, "Coloración", 120, 500.00)
        ]
        
        self.clientes = []
        self.citas = []
        self.next_cliente_id = 1
        self.next_cita_id = 1
        
        self.excel_file = "reservas_salon.xlsx"
        self._load_data()
    
    class Servicio:
        def __init__(self, id, nombre, duracion, precio):
            self.id = id
            self.nombre = nombre
            self.duracion = duracion
            self.precio = precio
    
    class Cliente:
        def __init__(self, id, nombre, email, telefono):
            self.id = id
            self.nombre = nombre
            self.email = email
            self.telefono = telefono
    
    class Cita:
        def __init__(self, id, cliente, servicio, fecha_hora):
            self.id = id
            self.cliente = cliente
            self.servicio = servicio
            self.fecha_hora = fecha_hora
            self.estado = "Confirmada"
        
        def get_details(self):
            return f"""
            Detalles de la cita:
            ID: {self.id}
            Cliente: {self.cliente.nombre}
            Servicio: {self.servicio.nombre}
            Fecha: {self.fecha_hora.strftime('%Y-%m-%d %I:%M %p')}
            Duración: {self.servicio.duracion} minutos
            Precio: ${self.servicio.precio:.2f}
            Estado: {self.estado}
            """
    
    def _load_data(self):
        if os.path.exists(self.excel_file):
            try:
                wb = load_workbook(self.excel_file)
                
                # Cargar clientes
                ws_clientes = wb["Clientes"]
                cliente_factory = ClienteFactory()
                for row in ws_clientes.iter_rows(min_row=2, values_only=True):
                    if row[0] is not None:
                        self.clientes.append(cliente_factory.create_entity(row[0], row[1], row[2], row[3]))
                        if row[0] >= self.next_cliente_id:
                            self.next_cliente_id = row[0] + 1
                
                # Cargar citas
                ws_citas = wb["Citas"]
                cita_factory = CitaFactory()
                for row in ws_citas.iter_rows(min_row=2, values_only=True):
                    if row[0] is not None:
                        cliente = next((c for c in self.clientes if c.id == row[1]), None)
                        servicio = next((s for s in self.servicios if s.id == row[2]), None)
                        if cliente and servicio:
                            # Convertir string de fecha a datetime (ahora acepta AM/PM)
                            fecha_str = row[3]
                            if isinstance(fecha_str, str):
                                fecha_hora = datetime.strptime(fecha_str, '%Y-%m-%d %I:%M %p')
                            else:
                                fecha_hora = row[3]
                            cita = cita_factory.create_entity(row[0], cliente, servicio, fecha_hora)
                            cita.estado = row[4]
                            self.citas.append(cita)
                            if row[0] >= self.next_cita_id:
                                self.next_cita_id = row[0] + 1
            except Exception as e:
                print(f"Error al cargar datos: {e}")
    
    def save_data(self):
        try:
            wb = Workbook()
            
            # Guardar clientes
            ws_clientes = wb.active
            ws_clientes.title = "Clientes"
            ws_clientes.append(["ID", "Nombre", "Email", "Teléfono"])
            for cliente in self.clientes:
                ws_clientes.append([cliente.id, cliente.nombre, cliente.email, cliente.telefono])
            
            # Guardar citas (ahora con formato AM/PM)
            ws_citas = wb.create_sheet("Citas")
            ws_citas.append(["ID", "ID Cliente", "Servicio", "Fecha", "Estado"])
            for cita in self.citas:
                ws_citas.append([
                    cita.id,
                    cita.cliente.id,
                    cita.servicio.nombre,
                    cita.fecha_hora.strftime('%Y-%m-%d %I:%M %p'),  # Guarda en AM/PM
                    cita.estado
                ])
            
            wb.save(self.excel_file)
        except Exception as e:
            print(f"Error al guardar datos: {e}")
    
    def agregar_cliente(self, nombre, email, telefono):
        cliente_factory = ClienteFactory()
        cliente = cliente_factory.create_entity(self.next_cliente_id, nombre, email, telefono)
        self.next_cliente_id += 1
        self.clientes.append(cliente)
        self.save_data()
        return cliente
    
    def agregar_cita(self, cliente, servicio, fecha_hora):
        cita_factory = CitaFactory()
        cita = cita_factory.create_entity(self.next_cita_id, cliente, servicio, fecha_hora)
        self.next_cita_id += 1
        self.citas.append(cita)
        self.save_data()
        return cita
    
    def cancelar_cita(self, cita_id):
        command = CancelarCitaCommand(self, cita_id)
        return command.execute()
    
    def obtener_citas_por_fecha(self, fecha):
        command = BuscarCitasCommand(self, fecha=fecha)
        success, result = command.execute()
        return result if success else []
    
    def obtener_citas_cliente(self, cliente_id):
        command = BuscarCitasCommand(self, cliente_id=cliente_id)
        success, result = command.execute()
        return result if success else []
    
    def verificar_disponibilidad(self, servicio, fecha_hora):
        hora_fin = fecha_hora + timedelta(minutes=servicio.duracion)
        
        for cita in self.citas:
            if cita.estado != "Confirmada":
                continue
                
            cita_fin = cita.fecha_hora + timedelta(minutes=cita.servicio.duracion)
            
            if (fecha_hora < cita_fin) and (hora_fin > cita.fecha_hora):
                return False
        
        return True

class SalonBellezaApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Bienvenido al Salon GuapasMakeup")
        self.root.geometry("800x600")
        self.root.configure(bg="#f8e1f4")
        
        # Estilos
        self.style = ttk.Style()
        self.style.configure("TFrame", background="#f8e1f4")
        self.style.configure("TLabel", background="#f8e1f4", font=("Arial", 10))
        self.style.configure("TButton", font=("Arial", 10), padding=5)
        self.style.configure("Header.TLabel", font=("Arial", 14, "bold"))
        
        # Inicializar base de datos
        self.db = Database()
        
        # Crear interfaz
        self.create_welcome_screen()
    
    def create_welcome_screen(self):
        self.clear_screen()
        
        header = ttk.Label(self.root, text="¡Bienvenido al Salon GuapasMakeup", style="Header.TLabel")
        header.pack(pady=20)
        
        # Botones principales
        btn_frame = ttk.Frame(self.root)
        btn_frame.pack(pady=20)
        
        ttk.Button(btn_frame, text="Ver Servicios", command=self.show_services).pack(pady=5, fill=tk.X)
        ttk.Button(btn_frame, text="Agendar Cita", command=self.create_booking_screen).pack(pady=5, fill=tk.X)
        ttk.Button(btn_frame, text="Mis Citas", command=self.show_my_bookings).pack(pady=5, fill=tk.X)
        ttk.Button(btn_frame, text="Modo Administrador", command=self.admin_login).pack(pady=5, fill=tk.X)
        ttk.Button(btn_frame, text="Salir", command=self.root.quit).pack(pady=5, fill=tk.X)
    
    def show_services(self):
        self.clear_screen()
        
        ttk.Label(self.root, text="Nuestros Servicios", style="Header.TLabel").pack(pady=10)
        
        services_frame = ttk.Frame(self.root)
        services_frame.pack(pady=10, padx=20, fill=tk.BOTH, expand=True)
        
        for service in self.db.servicios:
            service_frame = ttk.Frame(services_frame, borderwidth=1, relief="solid")
            service_frame.pack(pady=5, fill=tk.X)
            
            ttk.Label(service_frame, text=f"{service.nombre}", font=("Arial", 12, "bold")).pack(anchor="w")
            ttk.Label(service_frame, text=f"Duración: {service.duracion} min").pack(anchor="w")
            ttk.Label(service_frame, text=f"Precio: ${service.precio:.2f}").pack(anchor="w")
        
        ttk.Button(self.root, text="Volver", command=self.create_welcome_screen).pack(pady=10)

    def create_booking_screen(self):
        self.clear_screen()
        
        ttk.Label(self.root, text="Agendar Nueva Cita", style="Header.TLabel").pack(pady=10)
        
        form_frame = ttk.Frame(self.root)
        form_frame.pack(pady=10, padx=20, fill=tk.BOTH)
        
        # Datos del cliente
        ttk.Label(form_frame, text="Datos Personales", font=("Arial", 12, "bold")).grid(row=0, column=0, columnspan=2, pady=5, sticky="w")
        
        ttk.Label(form_frame, text="Nombre completo:").grid(row=1, column=0, sticky="e", padx=5, pady=5)
        self.nombre_entry = ttk.Entry(form_frame)
        self.nombre_entry.grid(row=1, column=1, sticky="ew", pady=5)
        
        ttk.Label(form_frame, text="Email:").grid(row=2, column=0, sticky="e", padx=5, pady=5)
        self.email_entry = ttk.Entry(form_frame)
        self.email_entry.grid(row=2, column=1, sticky="ew", pady=5)
        
        ttk.Label(form_frame, text="Teléfono:").grid(row=3, column=0, sticky="e", padx=5, pady=5)
        self.telefono_entry = ttk.Entry(form_frame)
        self.telefono_entry.grid(row=3, column=1, sticky="ew", pady=5)
        
        # Selección de servicio
        ttk.Label(form_frame, text="Servicio", font=("Arial", 12, "bold")).grid(row=4, column=0, columnspan=2, pady=5, sticky="w")
        
        ttk.Label(form_frame, text="Seleccione servicio:").grid(row=5, column=0, sticky="e", padx=5, pady=5)
        self.servicio_var = tk.StringVar()
        servicios = [f"{s.id}. {s.nombre} (${s.precio:.2f})" for s in self.db.servicios]
        self.servicio_combobox = ttk.Combobox(form_frame, textvariable=self.servicio_var, values=servicios, state="readonly")
        self.servicio_combobox.grid(row=5, column=1, sticky="ew", pady=5)
        
        # Fecha y hora (ahora con AM/PM)
        ttk.Label(form_frame, text="Fecha y Hora", font=("Arial", 12, "bold")).grid(row=6, column=0, columnspan=2, pady=5, sticky="w")
        
        ttk.Label(form_frame, text="Fecha (AAAA-MM-DD):").grid(row=7, column=0, sticky="e", padx=5, pady=5)
        self.fecha_entry = ttk.Entry(form_frame)
        self.fecha_entry.grid(row=7, column=1, sticky="ew", pady=5)
        self.fecha_entry.insert(0, datetime.now().strftime("%Y-%m-%d"))
        
        ttk.Label(form_frame, text="Hora (HH:MM AM/PM):").grid(row=8, column=0, sticky="e", padx=5, pady=5)
        self.hora_entry = ttk.Entry(form_frame)
        self.hora_entry.grid(row=8, column=1, sticky="ew", pady=5)
        self.hora_entry.insert(0, "02:00 PM")  # Ejemplo por defecto en AM/PM
        
        # Botones
        button_frame = ttk.Frame(form_frame)
        button_frame.grid(row=9, column=0, columnspan=2, pady=10)
        
        ttk.Button(button_frame, text="Agendar", command=self.process_booking).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="Cancelar", command=self.create_welcome_screen).pack(side=tk.LEFT, padx=5)
    
    def process_booking(self):
        # Validar datos
        nombre = self.nombre_entry.get()
        email = self.email_entry.get()
        telefono = self.telefono_entry.get()
        servicio_str = self.servicio_var.get()
        fecha_str = self.fecha_entry.get()
        hora_str = self.hora_entry.get()
        
        if not all([nombre, email, telefono, servicio_str, fecha_str, hora_str]):
            messagebox.showerror("Error", "Todos los campos son obligatorios")
            return
        
        try:
            servicio_id = int(servicio_str.split(".")[0])
            # Convertir a datetime con AM/PM
            fecha_hora = datetime.strptime(f"{fecha_str} {hora_str}", "%Y-%m-%d %I:%M %p")
        except ValueError:
            messagebox.showerror("Error", "Formato de fecha/hora o servicio inválido. Use formato: 'HH:MM AM/PM'")
            return
        
        # Registrar cliente
        cliente = self.db.agregar_cliente(nombre, email, telefono)
        
        # Agendar cita usando Command
        servicio = next(s for s in self.db.servicios if s.id == servicio_id)
        
        agendar_command = AgendarCitaCommand(self.db, cliente, servicio, fecha_hora)
        success, result = agendar_command.execute()
        
        if not success:
            messagebox.showerror("Error", result)
            return
        
        cita = result
        
        # Aplicar decoradores según condiciones
        if datetime.now().weekday() in [0, 1]:  # Lunes o martes
            cita = CitaConDescuento(cita, 10)  # 10% de descuento
        
        if "@" in cliente.email:  # Si tiene email, aplicar notificación
            cita = CitaConNotificacion(cita)
        
        # Mostrar confirmación con detalles decorados
        messagebox.showinfo("Confirmación", cita.get_details())
        self.create_welcome_screen()
    
    def show_my_bookings(self):
        self.clear_screen()
        
        ttk.Label(self.root, text="Mis Citas", style="Header.TLabel").pack(pady=10)
        
        search_frame = ttk.Frame(self.root)
        search_frame.pack(pady=10)
        
        ttk.Label(search_frame, text="Ingrese su email:").pack(side=tk.LEFT, padx=5)
        self.email_search_entry = ttk.Entry(search_frame, width=30)
        self.email_search_entry.pack(side=tk.LEFT, padx=5)
        ttk.Button(search_frame, text="Buscar", command=self.search_bookings).pack(side=tk.LEFT, padx=5)
        
        self.bookings_frame = ttk.Frame(self.root)
        self.bookings_frame.pack(pady=10, fill=tk.BOTH, expand=True)
        
        ttk.Button(self.root, text="Volver", command=self.create_welcome_screen).pack(pady=10)
    
    def search_bookings(self):
        email = self.email_search_entry.get()
        if not email:
            messagebox.showerror("Error", "Por favor ingrese su email")
            return
        
        cliente = next((c for c in self.db.clientes if c.email.lower() == email.lower()), None)
        if not cliente:
            messagebox.showerror("Error", "No se encontraron citas con ese email")
            return
        
        # Limpiar frame de resultados
        for widget in self.bookings_frame.winfo_children():
            widget.destroy()
        
        citas = self.db.obtener_citas_cliente(cliente.id)
        
        if not citas:
            ttk.Label(self.bookings_frame, text="No tiene citas agendadas").pack()
            return
        
        ttk.Label(self.bookings_frame, text=f"Citas para {cliente.nombre}:", font=("Arial", 12, "bold")).pack(anchor="w")
        
        for i, cita in enumerate(citas, 1):
            cita_frame = ttk.Frame(self.bookings_frame, borderwidth=1, relief="solid")
            cita_frame.pack(pady=5, fill=tk.X)
            
            ttk.Label(cita_frame, text=f"Cita #{i}", font=("Arial", 10, "bold")).grid(row=0, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Servicio: {cita.servicio.nombre}").grid(row=1, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Fecha: {cita.fecha_hora.strftime('%Y-%m-%d %I:%M %p')}").grid(row=2, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Estado: {cita.estado}").grid(row=3, column=0, sticky="w")
            
            if cita.estado == "Confirmada":
                ttk.Button(cita_frame, text="Cancelar", 
                          command=lambda c=cita: self.cancel_booking(c)).grid(row=0, column=1, rowspan=4, padx=5)
    
    def cancel_booking(self, cita):
        if messagebox.askyesno("Confirmar", "¿Está seguro que desea cancelar esta cita?"):
            command = CancelarCitaCommand(self.db, cita.id)
            success, result = command.execute()
            if success:
                messagebox.showinfo("Éxito", result)
                self.search_bookings()  # Refrescar la lista
            else:
                messagebox.showerror("Error", result)
    
    def admin_login(self):
        self.clear_screen()
        
        ttk.Label(self.root, text="Acceso Administrador", style="Header.TLabel").pack(pady=20)
        
        login_frame = ttk.Frame(self.root)
        login_frame.pack(pady=10)
        
        ttk.Label(login_frame, text="Contraseña:").grid(row=0, column=0, padx=5, pady=5)
        self.password_entry = ttk.Entry(login_frame, show="*")
        self.password_entry.grid(row=0, column=1, padx=5, pady=5)
        
        button_frame = ttk.Frame(login_frame)
        button_frame.grid(row=1, column=0, columnspan=2, pady=10)
        
        ttk.Button(button_frame, text="Ingresar", command=self.check_admin_password).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="Volver", command=self.create_welcome_screen).pack(side=tk.LEFT, padx=5)
    
    def check_admin_password(self):
        # Contraseña simple para demostración
        if self.password_entry.get() == "admin123":
            self.show_admin_panel()
        else:
            messagebox.showerror("Error", "Contraseña incorrecta")
    
    def show_admin_panel(self):
        self.clear_screen()
        
        ttk.Label(self.root, text="Panel de Administración", style="Header.TLabel").pack(pady=10)
        
        # Pestañas
        tab_control = ttk.Notebook(self.root)
        
        # Pestaña de citas
        tab_citas = ttk.Frame(tab_control)
        tab_control.add(tab_citas, text="Ver Citas")
        
        # Controles para seleccionar fecha
        date_frame = ttk.Frame(tab_citas)
        date_frame.pack(pady=10)
        
        ttk.Label(date_frame, text="Seleccione fecha:").pack(side=tk.LEFT, padx=5)
        self.admin_date_entry = ttk.Entry(date_frame)
        self.admin_date_entry.pack(side=tk.LEFT, padx=5)
        self.admin_date_entry.insert(0, datetime.now().strftime("%Y-%m-%d"))
        
        ttk.Button(date_frame, text="Buscar", command=self.show_admin_bookings).pack(side=tk.LEFT, padx=5)
        
        # Frame para resultados
        self.admin_results_frame = ttk.Frame(tab_citas)
        self.admin_results_frame.pack(fill=tk.BOTH, expand=True)
        
        tab_control.pack(expand=1, fill="both")
        
        ttk.Button(self.root, text="Volver", command=self.create_welcome_screen).pack(pady=10)
        
        # Mostrar citas del día actual
        self.show_admin_bookings()
    
    def show_admin_bookings(self):
        date_str = self.admin_date_entry.get()
        
        try:
            fecha = datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            messagebox.showerror("Error", "Formato de fecha inválido. Use AAAA-MM-DD")
            return
        
        # Limpiar frame de resultados
        for widget in self.admin_results_frame.winfo_children():
            widget.destroy()
        
        citas = self.db.obtener_citas_por_fecha(fecha)
        
        if not citas:
            ttk.Label(self.admin_results_frame, text=f"No hay citas para el {fecha.strftime('%Y-%m-%d')}").pack()
            return
        
        ttk.Label(self.admin_results_frame, text=f"Citas para el {fecha.strftime('%Y-%m-%d')}:", 
                 font=("Arial", 12, "bold")).pack(anchor="w")
        
        for cita in citas:
            cita_frame = ttk.Frame(self.admin_results_frame, borderwidth=1, relief="solid")
            cita_frame.pack(pady=5, fill=tk.X)
            
            ttk.Label(cita_frame, text=f"Cliente: {cita.cliente.nombre}").grid(row=0, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Servicio: {cita.servicio.nombre}").grid(row=1, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Hora: {cita.fecha_hora.strftime('%I:%M %p')}").grid(row=2, column=0, sticky="w")
            ttk.Label(cita_frame, text=f"Duración: {cita.servicio.duracion} min").grid(row=3, column=0, sticky="w")
    
    def clear_screen(self):
        for widget in self.root.winfo_children():
            widget.destroy()

if __name__ == "__main__":
    root = tk.Tk()
    app = SalonBellezaApp(root)
    root.mainloop()